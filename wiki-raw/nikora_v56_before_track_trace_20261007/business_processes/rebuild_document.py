from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
root=Path(__file__).resolve().parent
readme=(root/"README.md").read_text(encoding="utf-8")
ordered=re.findall(r"\| \d+ \| \[([^]]+)\]\(([^)]+\.md)\)",readme)
assert len(ordered)==15
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
sec.left_margin=Cm(2.3);sec.right_margin=Cm(2)
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2"]:
    style=doc.styles[name]
    style.font.name="Calibri";style.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles["Normal"]
normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
normal.paragraph_format.keep_together=True
for name,size in [("Title",22),("Heading 1",15),("Heading 2",12)]:
    s=doc.styles[name];s.font.size=Pt(size)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(12 if name!="Title" else 0)
    s.paragraph_format.space_after=Pt(8)
def add_markdown_block(block):
    block=block.strip()
    if not block or block.startswith("[Оглавление"):
        return
    if block.startswith("|"):
        lines=block.splitlines()
        rows=[[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
        assert all(re.fullmatch(r":?-+:?",cell) for cell in rows[1])
        rows=rows[:1]+rows[2:]
        table=doc.add_table(rows=len(rows),cols=len(rows[0]))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        widths=[2.5,8.3,5.9] if rows[0][0]=="Уровень" else [3.3,7.0,6.4]
        for col,width in zip(table.columns,widths): col.width=Cm(width)
        pr=table._tbl.tblPr
        borders=OxmlElement("w:tblBorders")
        for name in ["top","left","bottom","right","insideH","insideV"]:
            border=OxmlElement("w:"+name)
            for k,v in {"val":"single","sz":"4","color":"D9D9D9"}.items(): border.set(qn("w:"+k),v)
            borders.append(border)
        pr.append(borders)
        for ri,(row,values) in enumerate(zip(table.rows,rows)):
            trpr=row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if ri==0: trpr.append(OxmlElement("w:tblHeader"))
            for ci,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement("w:tcMar")
                for name in ["top","left","bottom","right"]:
                    e=OxmlElement("w:"+name);e.set(qn("w:w"),"100");e.set(qn("w:type"),"dxa");margins.append(e)
                tcpr.append(margins)
                if ri==0:
                    shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"DCE6F1");tcpr.append(shade)
                par=cell.paragraphs[0]
                par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.space_before=Pt(0)
                par.paragraph_format.line_spacing=1.05
                par.paragraph_format.keep_with_next=(ri==0)
                run=par.add_run(value);run.font.size=Pt(10.5);run.bold=(ri==0)
                if rows[0][0]=="Уровень" and ci==0: par.alignment=1
        doc.add_paragraph().paragraph_format.space_after=Pt(3)
    elif block.startswith("#"):
        doc.add_paragraph(block.lstrip("#").strip(),"Heading 2")
    elif block.startswith("- "):
        for item in block.splitlines(): doc.add_paragraph(item[2:],"List Bullet")
    else:
        doc.add_paragraph(block)

doc.core_properties.title="Процессы снабжения магазинов Никора"
doc.core_properties.subject="Транспортный план и организация комплектации по волнам"
doc.core_properties.author=""
doc.core_properties.keywords="Никора, транспортное планирование, волны, пополнение"
doc.add_paragraph("Процессы снабжения магазинов Никора","Title")
doc.add_paragraph("Редакция 55 от 7 октября 2026 года")
for paragraph in readme.split("\n\n"):
    if paragraph.startswith("Рабочий цикл") or paragraph.startswith("Последовательность работы"):
        doc.add_paragraph(paragraph)
doc.add_paragraph("Термины","Heading 2")
section=readme.split("## Термины")[1].split("## Описания процессов")[0].strip()
for p in section.split("\n\n"):
    doc.add_paragraph(p)
doc.add_paragraph("Содержание","Heading 2")
for i,(title,name) in enumerate(ordered,1):
    p=doc.add_paragraph(f"{i}. {title}")
    p.paragraph_format.space_after=Pt(3)
doc.add_paragraph("Состав паллет и динамические диапазоны отбора")
doc.add_paragraph("Управление ресурсами и назначение заданий WMS")
doc.add_page_break()
for i,(title,name) in enumerate(ordered,1):
    doc.add_paragraph(f"{i} {title}","Heading 1")
    for block in (root/name).read_text(encoding="utf-8").split("\n\n")[1:]:
        add_markdown_block(block)
slotting=(root/"dynamic_slotting.md").read_text(encoding="utf-8")
doc.add_paragraph("Формирование состава паллет и управление динамическими диапазонами отбора","Heading 1")
for block in slotting.split("\n\n")[1:]:
    add_markdown_block(block)
resource=(root/"resource_management.md").read_text(encoding="utf-8")
doc.add_paragraph("Управление ресурсами и назначение заданий WMS","Heading 1")
for block in resource.split("\n\n")[1:]:
    add_markdown_block(block)
footer=sec.footer.paragraphs[0];footer.alignment=2
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
footer.add_run()._r.append(field)
for element in list(doc.styles.element.iter(qn("w:pBdr"))):
    element.getparent().remove(element)
out=root/"Nikora_business_processes_v55.docx"
doc.save(out)
text="\n".join([p.text for p in doc.paragraphs]+[cell.text for table in doc.tables for row in table.rows for cell in row.cells])
for name in [name for _,name in ordered]+["dynamic_slotting.md","resource_management.md"]:
    for block in (root/name).read_text(encoding="utf-8").split("\n\n")[1:]:
        block=block.strip()
        if not block or block.startswith("[Оглавление"): continue
        if block.startswith("|"):
            rows=[[c.strip() for c in line.strip().strip("|").split("|")] for line in block.splitlines()]
            assert all(cell in text for row in rows[:1]+rows[2:] for cell in row),name
        elif block.startswith("#"): assert block.lstrip("#").strip() in text,name
        elif block.startswith("- "): assert all(item[2:] in text for item in block.splitlines()),name
        else: assert block in text,name
assert len(doc.tables)==2
assert ordered[0][1]=="06_transport_planning.md"
print("Создан документ; 15 процессов, динамические диапазоны и стратегия ресурсов WMS, 2 таблицы; полнота текста проверена")
