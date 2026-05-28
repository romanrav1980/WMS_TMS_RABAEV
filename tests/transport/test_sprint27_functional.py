"""
test_sprint27_functional.py — Functional tests for Sprint 27.

Sprint 27 scope:
  - GET /billing/orders/export.xlsx — export filtered billing registry to Excel
  - Verifies XLSX content: header row, order rows, totals

Run:
    pytest tests/transport/test_sprint27_functional.py -v
"""

import pytest
import requests

from tests.transport.config import API_BASE as BASE, API_AUTH as AUTH

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_MAGIC = b"PK\x03\x04"


def get(path: str, **kwargs) -> requests.Response:
    return requests.get(f"{BASE}{path}", auth=AUTH, **kwargs)


# ---------------------------------------------------------------------------
# Registry export endpoint
# ---------------------------------------------------------------------------

class TestRegistryExportEndpoint:

    def test_export_endpoint_reachable(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx")
        assert r.status_code in (200, 401, 403)

    def test_export_returns_xlsx_mime(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        ct = r.headers.get("content-type", "")
        assert XLSX_MIME in ct, f"Expected XLSX content-type, got: {ct}"

    def test_export_is_valid_zip(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC

    def test_export_content_disposition(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower()
        assert ".xlsx" in cd.lower()

    def test_export_with_company_filter(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx?company=TEST_MISSING_CO_XYZ")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC

    def test_export_with_date_filter(self):
        r = get("/api/admin/transport/billing/orders/export.xlsx?date_from=2026-01-01&date_to=2026-12-31")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC

    def test_export_does_not_conflict_with_order_id_routes(self):
        """Ensure /billing/orders/export.xlsx does not match /{order_id} route."""
        r = get("/api/admin/transport/billing/orders/export.xlsx")
        # Should NOT return 422 (which would mean it tried to parse "export.xlsx" as int)
        assert r.status_code != 422, "Route conflict: server tried to parse 'export.xlsx' as integer"


# ---------------------------------------------------------------------------
# XLSX content validation
# ---------------------------------------------------------------------------

class TestRegistryExportContent:

    def _get_workbook(self, params: str = ""):
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed on test runner")
        r = get(f"/api/admin/transport/billing/orders/export.xlsx{params}")
        if r.status_code != 200:
            pytest.skip("Export not accessible")
        return load_workbook(io.BytesIO(r.content))

    def test_workbook_has_one_sheet(self):
        wb = self._get_workbook()
        assert len(wb.sheetnames) >= 1

    def test_sheet_has_header_rows(self):
        wb = self._get_workbook()
        ws = wb.active
        assert ws.max_row >= 4, f"Expected at least 4 rows (title+filter+blank+header)"

    def test_header_row_contains_column_labels(self):
        wb = self._get_workbook()
        ws = wb.active
        # Row 4 should have column headers
        row4 = [str(ws.cell(4, c).value or "").lower() for c in range(1, 10)]
        combined = " ".join(row4)
        assert any(kw in combined for kw in ["счёт", "счет", "num", "компания", "рейс"]), \
            f"Row 4 doesn't look like header: {row4}"

    def test_row_count_matches_api(self):
        """Number of data rows in Excel equals number of orders from API."""
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed")
        orders_r = get("/api/admin/transport/billing/orders")
        if orders_r.status_code != 200:
            pytest.skip("Orders endpoint unavailable")
        api_count = len(orders_r.json())
        xlsx_r = get("/api/admin/transport/billing/orders/export.xlsx")
        if xlsx_r.status_code != 200:
            pytest.skip("Export not accessible")
        wb = load_workbook(io.BytesIO(xlsx_r.content))
        ws = wb.active
        # data rows = max_row minus 4 (title + filter + blank + header) minus 1 (total)
        data_rows = ws.max_row - 4 - 1
        assert data_rows == api_count, \
            f"XLSX has {data_rows} data rows but API returned {api_count} orders"

    def test_filtered_export_has_fewer_rows(self):
        """Export with status=open should have <= unfiltered count."""
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed")
        all_r = get("/api/admin/transport/billing/orders/export.xlsx")
        open_r = get("/api/admin/transport/billing/orders/export.xlsx?closed=0&payed=0")
        if all_r.status_code != 200 or open_r.status_code != 200:
            pytest.skip("Export not accessible")
        wb_all = load_workbook(io.BytesIO(all_r.content))
        wb_open = load_workbook(io.BytesIO(open_r.content))
        all_rows = wb_all.active.max_row
        open_rows = wb_open.active.max_row
        assert open_rows <= all_rows, "Filtered export has more rows than unfiltered"


# ---------------------------------------------------------------------------
# Regression: per-order export still works
# ---------------------------------------------------------------------------

class TestPerOrderExportRegression:

    def test_per_order_export_still_returns_xlsx(self):
        r = get("/api/admin/transport/billing/orders")
        if r.status_code != 200:
            pytest.skip("No access")
        orders = r.json()
        if not orders:
            pytest.skip("No orders")
        oid = orders[0]["order_id"]
        r2 = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        assert r2.status_code in (200, 401, 403, 404)
        if r2.status_code == 200:
            assert r2.content[:4] == XLSX_MAGIC
