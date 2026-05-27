"""
test_sprint28_functional.py — Functional tests for Sprint 28.

Sprint 28 scope:
  - GET /tasks/export.xlsx — export trips list to Excel (TZ §3 «В Excel»)
  - Supports same filters as GET /tasks: shipment_date, condition, task_id, etc.

Run:
    pytest tests/transport/test_sprint28_functional.py -v
"""

import pytest
import requests
from datetime import date

BASE = "http://127.0.0.1:8088"
AUTH = ("admin", "admin123")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_MAGIC = b"PK\x03\x04"
TODAY = date.today().isoformat()


def get(path: str, **kwargs) -> requests.Response:
    return requests.get(f"{BASE}{path}", auth=AUTH, **kwargs)


# ---------------------------------------------------------------------------
# Endpoint existence and basic response
# ---------------------------------------------------------------------------

class TestTasksExportEndpoint:

    def test_endpoint_reachable(self):
        r = get("/api/admin/transport/tasks/export.xlsx")
        assert r.status_code in (200, 401, 403)

    def test_returns_xlsx_mime(self):
        r = get("/api/admin/transport/tasks/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        ct = r.headers.get("content-type", "")
        assert XLSX_MIME in ct, f"Got content-type: {ct}"

    def test_returns_valid_zip_magic(self):
        r = get("/api/admin/transport/tasks/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC

    def test_content_disposition_attachment(self):
        r = get("/api/admin/transport/tasks/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower()
        assert ".xlsx" in cd.lower()

    def test_filename_contains_date(self):
        r = get(f"/api/admin/transport/tasks/export.xlsx?shipment_date={TODAY}")
        if r.status_code != 200:
            pytest.skip("No access")
        cd = r.headers.get("content-disposition", "")
        assert TODAY in cd, f"Expected date in filename, got: {cd}"

    def test_no_route_conflict_with_task_id(self):
        """Ensure /tasks/export.xlsx is not misinterpreted as /tasks/{task_id}."""
        r = get("/api/admin/transport/tasks/export.xlsx")
        assert r.status_code != 422, "Route conflict: server tried to parse 'export.xlsx' as task_id"

    def test_with_date_filter(self):
        r = get(f"/api/admin/transport/tasks/export.xlsx?shipment_date={TODAY}")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC

    def test_with_condition_filter(self):
        r = get("/api/admin/transport/tasks/export.xlsx?condition=Отгружен")
        if r.status_code != 200:
            pytest.skip("No access")
        assert r.content[:4] == XLSX_MAGIC


# ---------------------------------------------------------------------------
# XLSX content validation
# ---------------------------------------------------------------------------

class TestTasksExportContent:

    def _load_wb(self, params: str = ""):
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed")
        r = get(f"/api/admin/transport/tasks/export.xlsx{params}")
        if r.status_code != 200:
            pytest.skip("Export not accessible")
        return load_workbook(io.BytesIO(r.content))

    def test_workbook_has_trips_sheet(self):
        wb = self._load_wb()
        assert len(wb.sheetnames) >= 1

    def test_header_row_present(self):
        wb = self._load_wb()
        ws = wb.active
        # Row 3 is the header row (row 1 = title, row 2 = empty/meta)
        assert ws.max_row >= 3

    def test_header_contains_trip_columns(self):
        wb = self._load_wb()
        ws = wb.active
        row3 = [str(ws.cell(3, c).value or "").lower() for c in range(1, 15)]
        combined = " ".join(row3)
        assert any(kw in combined for kw in ["дата", "пал", "машина", "статус"]), \
            f"Row 3 doesn't look like trip header: {row3}"

    def test_data_rows_match_api(self):
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed")
        params = f"?shipment_date={TODAY}"
        api_r = get(f"/api/admin/transport/tasks{params}")
        if api_r.status_code != 200:
            pytest.skip("Tasks API not accessible")
        api_count = len(api_r.json())
        xlsx_r = get(f"/api/admin/transport/tasks/export.xlsx{params}")
        if xlsx_r.status_code != 200:
            pytest.skip("Export not accessible")
        wb = load_workbook(io.BytesIO(xlsx_r.content))
        ws = wb.active
        # data rows = max_row - 3 (title + blank + header)
        data_rows = ws.max_row - 3
        assert data_rows == api_count, \
            f"XLSX has {data_rows} rows but API returned {api_count} tasks"


# ---------------------------------------------------------------------------
# Regression: tasks list endpoint unaffected
# ---------------------------------------------------------------------------

class TestTasksListRegression:

    def test_tasks_list_still_works(self):
        r = get(f"/api/admin/transport/tasks?shipment_date={TODAY}")
        assert r.status_code in (200, 401, 403)
        if r.status_code == 200:
            assert isinstance(r.json(), list)
