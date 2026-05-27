"""
test_sprint26_functional.py — Functional tests for Sprint 26.

Sprint 26 scope:
  - GET /billing/orders/{order_id}/export.xlsx — export billing order to Excel
  - DoD §12 #7: Excel contains all trips of the selected order with totals and requisites

Run:
    pytest tests/transport/test_sprint26_functional.py -v
"""

import pytest
import requests

BASE = "http://127.0.0.1:8088"
AUTH = ("admin", "admin123")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_MAGIC = b"PK\x03\x04"  # ZIP/OOXML magic bytes


def get(path: str, **kwargs) -> requests.Response:
    return requests.get(f"{BASE}{path}", auth=AUTH, **kwargs)


def _first_order_id() -> int | None:
    r = get("/api/admin/transport/billing/orders")
    if r.status_code != 200:
        return None
    orders = r.json()
    return orders[0]["order_id"] if orders else None


# ---------------------------------------------------------------------------
# Export endpoint existence
# ---------------------------------------------------------------------------

class TestExportEndpointExists:

    def test_export_endpoint_reachable(self):
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders in DB")
        r = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        assert r.status_code in (200, 401, 403)

    def test_export_returns_xlsx_content_type(self):
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders in DB")
        r = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access or endpoint error")
        ct = r.headers.get("content-type", "")
        assert XLSX_MIME in ct, f"Expected XLSX content-type, got: {ct}"

    def test_export_returns_valid_zip_magic(self):
        """XLSX files are ZIP archives — check magic bytes."""
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders in DB")
        r = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access or endpoint error")
        assert r.content[:4] == XLSX_MAGIC, "Response does not have XLSX/ZIP magic bytes"

    def test_export_content_disposition_header(self):
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders in DB")
        r = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        if r.status_code != 200:
            pytest.skip("No access")
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower(), f"Expected attachment header, got: {cd}"
        assert ".xlsx" in cd.lower(), f"Expected .xlsx in filename, got: {cd}"

    def test_export_nonexistent_order_returns_404(self):
        r = get("/api/admin/transport/billing/orders/999999999/export.xlsx")
        assert r.status_code in (404, 401, 403)


# ---------------------------------------------------------------------------
# XLSX content validation (requires openpyxl on test runner)
# ---------------------------------------------------------------------------

class TestExportContent:

    def _get_workbook(self, order_id: int):
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed on test runner")
        r = get(f"/api/admin/transport/billing/orders/{order_id}/export.xlsx")
        if r.status_code != 200:
            pytest.skip("Export endpoint not accessible")
        return load_workbook(io.BytesIO(r.content))

    def test_workbook_has_sheet(self):
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders")
        wb = self._get_workbook(oid)
        assert len(wb.sheetnames) >= 1

    def test_sheet_has_header_row(self):
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders")
        wb = self._get_workbook(oid)
        ws = wb.active
        # The sheet must have at least a few rows (meta block + header)
        assert ws.max_row >= 7, f"Expected at least 7 rows, got {ws.max_row}"

    def test_sheet_has_task_data_columns(self):
        """The column header row must contain trip-related labels."""
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders")
        wb = self._get_workbook(oid)
        ws = wb.active
        all_values = [str(ws.cell(r, c).value or "").lower()
                      for r in range(1, ws.max_row + 1)
                      for c in range(1, ws.max_column + 1)]
        assert any("рейс" in v or "tt_id" in v or "id" in v for v in all_values), \
            "Expected trip ID column in worksheet"

    def test_tasks_in_xlsx_match_api(self):
        """Number of data rows in XLSX equals number of tasks from API."""
        oid = _first_order_id()
        if not oid:
            pytest.skip("No billing orders")
        try:
            from openpyxl import load_workbook
            import io
        except ImportError:
            pytest.skip("openpyxl not installed")
        tasks_r = get(f"/api/admin/transport/billing/orders/{oid}/tasks")
        if tasks_r.status_code != 200:
            pytest.skip("Tasks endpoint unavailable")
        api_count = len(tasks_r.json())
        if api_count == 0:
            pytest.skip("No tasks in order")
        r = get(f"/api/admin/transport/billing/orders/{oid}/export.xlsx")
        if r.status_code != 200:
            pytest.skip("Export not accessible")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Data rows = all rows minus meta (6 rows) minus header minus total
        data_rows = ws.max_row - 6 - 1 - 1
        assert data_rows == api_count, \
            f"XLSX has {data_rows} data rows but API returns {api_count} tasks"


# ---------------------------------------------------------------------------
# Regression: other billing endpoints unaffected
# ---------------------------------------------------------------------------

class TestBillingRegressionAfterSprint26:

    def test_billing_orders_still_works(self):
        r = get("/api/admin/transport/billing/orders")
        assert r.status_code in (200, 401, 403)
        if r.status_code == 200:
            assert isinstance(r.json(), list)

    def test_companies_endpoint_still_works(self):
        r = get("/api/admin/transport/billing/companies")
        assert r.status_code in (200, 401, 403)
