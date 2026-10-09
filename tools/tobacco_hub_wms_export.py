"""Reuse the checked DOCX styling without modifying the retail package."""
from pathlib import Path
import json
from docx import Document
import retail_standard_export as base

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'wiki/requirements/tobacco_regional_hubs'
QA=ROOT/'tmp/tobacco-hub-wms-qa'

def main():
    base.SRC=SRC
    base.OUT=SRC/'docx'
    base.QA=QA
    base.OUT.mkdir(exist_ok=True)
    QA.mkdir(parents=True,exist_ok=True)
    files=[p for p in sorted(SRC.glob('*.md')) if p.name!='publication.md']
    assert len(files)==6
    records=[]
    for p in files:
        records.append(base.make(p))
        dest=base.OUT/(p.stem+'.docx')
        d=Document(dest)
        for sec in d.sections:
            for para in sec.footer.paragraphs:
                for run in para.runs:
                    if 'Retail WMS TMS' in run.text:run.text=run.text.replace('Retail WMS TMS','WMS Tobacco Hubs')
        d.core_properties.author='Требования к WMS региональных хабов'
        d.core_properties.subject='Табачная продукция, марки, несколько собственников и SAP'
        for rel in d.part.rels.values():
            if rel.reltype==base.RT.HYPERLINK and rel.target_ref=='publication.docx':
                rel._target='../publication.md'
        if p.stem=='03_vendor_questionnaire':
            active=False
            for para in d.paragraphs:
                if para.style.name.startswith('Heading'):
                    active=para.text.startswith('TH ')
                if active:
                    para.paragraph_format.keep_with_next=not para.text.startswith('Проверяющий,')
        d.save(dest)
    (QA/'docx_structure.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(records,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
