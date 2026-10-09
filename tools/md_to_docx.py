from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SRC=Path(r"C:\projects\TMS\wiki\requirements\nikora_retail_selection"); OUT=SRC/"docx"; OUT.mkdir(exist_ok=True)
def shade(c,f):
    p=c._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:fill"),f); p.append(s)
def borders(t):
    p=t._tbl.tblPr; b=OxmlElement("w:tblBorders")
    for e in ("top","left","bottom","right","insideH","insideV"):
        x=OxmlElement("w:"+e); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"4"); x.set(qn("w:color"),"D9D9D9"); b.append(x)
    p.append(b)
def clean(s):
    s=re.sub(r"\[([^]]+)\]\(([^)]+)\)",r"\1",s); s=re.sub(r"\[([^]]+)\]\[[^]]+\]",r"\1",s); s=re.sub(r"\*\*([^*]+)\*\*",r"\1",s); return re.sub(r"\x60([^\x60]+)\x60",r"\1",s).strip()
def make(src):
    lines=src.read_text(encoding="utf-8").splitlines(); d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    st=d.styles; st["Normal"].font.name="Aptos"; st["Normal"].font.size=Pt(9); st["Normal"].paragraph_format.space_after=Pt(4)
    for n,z in [("Title",20),("Heading 1",15),("Heading 2",12),("Heading 3",10)]: st[n].font.name="Aptos Display"; st[n].font.size=Pt(z); st[n].font.color.rgb=RGBColor(0,0,0)
    i=0
    while i<len(lines):
        l=lines[i]
        if not l.strip() or l.startswith("<!--"): i+=1; continue
        if l.startswith("# "): p=d.add_paragraph(style="Title"); p.add_run(clean(l[2:])).bold=True; i+=1; continue
        m=re.match(r"^(#{2,3})\s+(.+)",l)
        if m: d.add_paragraph(clean(m.group(2)),style="Heading "+str(len(m.group(1))-1)); i+=1; continue
        if l.startswith("|") and i+1<len(lines) and "---" in lines[i+1]:
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*:?-{3,}",lines[i]): rows.append([clean(x) for x in lines[i].strip().strip("|").split("|")])
                i+=1
            if rows:
                t=d.add_table(rows=0,cols=max(map(len,rows))); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
                for ri,row in enumerate(rows):
                    cs=t.add_row().cells
                    for ci,c in enumerate(cs):
                        c.text=row[ci] if ci<len(row) else ""; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if ri==0: shade(c,"1F4E78")
                        elif ri%2==0: shade(c,"F3F6F8")
                        for r in c.paragraphs[0].runs: r.font.size=Pt(7.5); r.bold=ri==0; r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
                d.add_paragraph(); continue
        if re.match(r"^[-*]\s+",l): d.add_paragraph(clean(re.sub(r"^[-*]\s+","",l)),style="List Bullet"); i+=1; continue
        if re.match(r"^\d+\.\s+",l): d.add_paragraph(clean(re.sub(r"^\d+\.\s+","",l)),style="List Number"); i+=1; continue
        d.add_paragraph(clean(l)); i+=1
    d.core_properties.title=clean(next((x[2:] for x in lines if x.startswith("# ")),src.stem)); d.core_properties.author="NIKORA requirements working group"; out=OUT/(src.stem+".docx"); d.save(out); return out
if __name__=="__main__":
    for p in sorted(SRC.glob("*.md")):
        if p.name!="publication.md": print(make(p))
