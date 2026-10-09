"""Scoped retail Word export. Windows fallback: managed artifact runtime unavailable."""
from pathlib import Path
import json,re,zipfile
from urllib.parse import unquote
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/"wiki/requirements/retail_convenience_standard"
OUT=SRC/"docx"
QA=ROOT/"tmp/retail-standard-qa"
TOKEN=re.compile(r"\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|\x60([^\x60]+)\x60")
def heading_text(text):
    return re.sub(r"\s+"," ",re.sub(r"[^\w\s]"," ",text,flags=re.UNICODE)).strip()
def inline(p,text,bold=False,size=None,color=None):
    def run(s):
        r=p.add_run(s);r.bold=bold;r.font.name="Arial";r.font.size=Pt(size or 11)
        r.font.complex_script=False
        rf=r._element.get_or_add_rPr().rFonts
        rf.set(qn("w:cs"),"Arial");rf.set(qn("w:eastAsia"),"Arial")
        if color:r.font.color.rgb=RGBColor.from_string(color)
        return r
    pos=0
    for m in TOKEN.finditer(text):
        if m.start()>pos:run(text[pos:m.start()])
        if m.group(1):
            target=unquote(m.group(2))
            if not re.match(r"^[a-z]+:",target):
                local=(SRC/target.split("#")[0]).resolve()
                if local.parent==SRC and local.suffix==".md":target=local.stem+".docx"
                else:target=local.as_uri()+("#"+target.split("#",1)[1] if "#" in target else "")
            rel=p.part.relate_to(target,RT.HYPERLINK,is_external=True)
            h=OxmlElement("w:hyperlink");h.set(qn("r:id"),rel)
            r=OxmlElement("w:r");prop=OxmlElement("w:rPr")
            rf=OxmlElement("w:rFonts")
            for attr in ["ascii","hAnsi","cs","eastAsia"]:rf.set(qn("w:"+attr),"Arial")
            prop.append(rf)
            c=OxmlElement("w:color");c.set(qn("w:val"),color or "215E86");prop.append(c)
            u=OxmlElement("w:u");u.set(qn("w:val"),"single");prop.append(u)
            if size:
                z=OxmlElement("w:sz");z.set(qn("w:val"),str(int(size*2)));prop.append(z)
            r.append(prop);t=OxmlElement("w:t");t.text=m.group(1);r.append(t);h.append(r);p._p.append(h)
        elif m.group(3):run(m.group(3)).bold=True
        else:run(m.group(4)).font.name="Consolas"
        pos=m.end()
    if pos<len(text):run(text[pos:])
def page(sec,landscape=False):
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width=Mm(297 if landscape else 210);sec.page_height=Mm(210 if landscape else 297)
    sec.left_margin=sec.right_margin=Mm(16)
    sec.top_margin=Mm(16);sec.bottom_margin=Mm(17)
    sec.footer_distance=Mm(8);sec.header_distance=Mm(7)
def table(doc,rows):
    n=len(rows[0]);wide=n>6
    if wide:page(doc.add_section(WD_SECTION_START.NEW_PAGE),True)
    width=265 if wide else 178
    if n==12:widths=[14,61,10]+[20]*9
    elif n==4:widths=[width*.19,width*.25,width*.28,width*.28]
    elif n==3:widths=[width*.19,width*.32,width*.49]
    elif n==2:widths=[width*.29,width*.71]
    else:widths=[width/n]*n
    scale=width/sum(widths);widths=[x*scale for x in widths]
    t=doc.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    pr=t._tbl.tblPr;borders=OxmlElement("w:tblBorders")
    for edge in ["top","left","bottom","right","insideH","insideV"]:
        b=OxmlElement("w:"+edge)
        for k,v in [("val","single"),("sz","4"),("color","D9D9D9")]:b.set(qn("w:"+k),v)
        borders.append(b)
    pr.append(borders);margins=OxmlElement("w:tblCellMar")
    vertical_padding="70" if wide and rows[1][0].startswith("RT") else "90"
    for k,v in [("top",vertical_padding),("bottom",vertical_padding),("left","90"),("right","90")]:
        el=OxmlElement("w:"+k);el.set(qn("w:w"),v);el.set(qn("w:type"),"dxa");margins.append(el)
    pr.append(margins)
    for ri,data in enumerate(rows):
        row=t.add_row();rp=row._tr.get_or_add_trPr();rp.append(OxmlElement("w:cantSplit"))
        if ri==0:rp.append(OxmlElement("w:tblHeader"))
        for ci,cell in enumerate(row.cells):
            cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"24445C" if ri==0 else ("F1F4F6" if ri%2==0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=ri==0
            if (wide and ci!=1) or (ci==0 and len(data[ci])<9):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,data[ci],ri==0,9.5 if wide else 10,"FFFFFF" if ri==0 else "000000")
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    if wide:page(doc.add_section(WD_SECTION_START.NEW_PAGE))
def make(path):
    d=Document();page(d.sections[0]);st=d.styles
    for name,size in [("Normal",11),("Title",23),("Heading 1",15),("Heading 2",12),("Heading 3",11)]:
        s=st[name];s.font.name="Arial";s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.underline=False
        s.font.bold=name!="Normal"
        rf=s.element.get_or_add_rPr().rFonts
        rf.set(qn("w:cs"),"Arial");rf.set(qn("w:eastAsia"),"Arial")
        for b in list(s.element.xpath("./w:pPr/w:pBdr")):b.getparent().remove(b)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.08
        if name!="Normal":s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(11)
    st["Title"].paragraph_format.space_before=Pt(0);st["Normal"].paragraph_format.widow_control=True
    foot=d.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=foot.add_run("Retail WMS TMS  |  ");r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    f=OxmlElement("w:fldSimple");f.set(qn("w:instr"),"PAGE");foot._p.append(f)
    lines=path.read_text(encoding="utf-8").splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line or line.startswith("<!--"):i+=1;continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"^\|\s*:?-+",lines[i+1]):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                l=lines[i].strip()
                if not re.match(r"^\|\s*:?-+",l):rows.append([x.strip() for x in l.strip("|").split("|")])
                i+=1
            if any(len(x)!=len(rows[0]) for x in rows):raise ValueError(f"Uneven table {path}")
            table(d,rows);continue
        h=re.match(r"^(#{1,4})\s+(.*)",line)
        if h:
            level=len(h[1]);d.add_paragraph(heading_text(h[2]),style="Title" if level==1 else f"Heading {level-1}")
        else:
            number=re.match(r"^\d+\.\s+(.*)",line);bullet=re.match(r"^[-*]\s+(.*)",line)
            p=d.add_paragraph(style="List Number" if number else "List Bullet" if bullet else "Normal")
            txt=number[1] if number else bullet[1] if bullet else line;inline(p,txt)
            if txt.startswith(("Профиль ","Вес ","Статус: ")):p.paragraph_format.keep_with_next=True
            if "________________" in txt:p.paragraph_format.space_after=Pt(7)
            if path.stem in ("04_questionnaire_wms","05_questionnaire_tms") and txt.startswith(("Как решение", "Покажите", "Код ответа", "Модуль и версия", "Настройка или")):
                p.paragraph_format.keep_with_next=True
        i+=1
    d.core_properties.title=heading_text(next(l[2:] for l in lines if l.startswith("# ")))
    d.core_properties.author="Исследование требований WMS и TMS"
    d.core_properties.subject="Общий стандарт для сети магазинов у дома"
    dest=OUT/(path.stem+".docx");d.save(dest)
    with zipfile.ZipFile(dest) as z:
        if z.testzip():raise ValueError("Broken DOCX")
    return {"file":dest.name,"bytes":dest.stat().st_size,"tables":len(d.tables),"paragraphs":len(d.paragraphs),"hyperlinks":sum(1 for r in d.part.rels.values() if r.reltype==RT.HYPERLINK)}
def main():
    OUT.mkdir(exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
    files=[p for p in sorted(SRC.glob("*.md")) if p.name!="publication.md"]
    assert len(files)==11,len(files)
    records=[make(p) for p in files]
    (QA/"docx_structure.json").write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(records,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
