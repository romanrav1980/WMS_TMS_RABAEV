"""Print-first, standalone supplier questionnaire; preserve the version 1 package."""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import retail_standard_export as base

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'wiki/requirements/tobacco_regional_hubs/supplier_questionnaire_v2'
QA=ROOT/'tmp/tobacco-questionnaire-v2-qa'

def table(doc,rows):
    n=len(rows[0])
    widths={2:[60,118],4:[97,24,29,28],6:[12,31,71,12,14,38]}[n]
    t=doc.add_table(rows=0,cols=n)
    t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for k,v in [('top','100'),('bottom','100'),('left','70'),('right','70')]:
        el=OxmlElement('w:'+k);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');margins.append(el)
    t._tbl.tblPr.append(margins)
    for ri,data in enumerate(rows):
        row=t.add_row();rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if ri==0:rp.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if ri==0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=ri==0
            if (n==6 and ci in (0,3,4)) or (n==4 and ci>0):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            value=data[ci]
            if n==4 and ri==0:value=['Предметная область','Вес','Балл','Не проверено'][ci]
            if n==6 and ri>0 and ci==5:value='Способ: ____\n____________\n____________'
            base.inline(p,value,ri==0,10.5,'000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    data=json.loads((ROOT/'tools/tobacco_questionnaire_v2_data.json').read_text(encoding='utf-8'))
    assert len(data['rows'])==100 and sum(r['weight'] for r in data['rows'])==100
    for code,_,weight in data['groups']:assert sum(r['weight'] for r in data['rows'] if r['group']==code)==weight
    base.SRC=SRC;base.OUT=SRC;base.table=table
    QA.mkdir(parents=True,exist_ok=True)
    base.make(SRC/'questionnaire.md')
    dest=SRC/'questionnaire.docx';d=Document(dest)
    for p in list(d.paragraphs):
        if p.text=='Вопросы по предметным областям':p._element.getparent().remove(p._element);continue
        if p.style.name=='Heading 2' and re.match(r'^\d+ ',p.text):
            p.paragraph_format.page_break_before=True
            p.style=d.styles['Heading 1']
        if p.text=='Сводная оценка':p.paragraph_format.page_break_before=True
        if p.text=='Как рассчитать результат':p.paragraph_format.page_break_before=True
    for sec in d.sections:
        for p in sec.footer.paragraphs:
            for run in p.runs:
                if 'Retail WMS TMS' in run.text:run.text='Опросник WMS  2.0  |  '
    d.core_properties.title='Опросник поставщика WMS для региональных табачных хабов'
    d.core_properties.author='Заказчик WMS'
    d.core_properties.subject='Самостоятельный табличный опросник с весами'
    d.save(dest)
    assert not any(r.reltype==base.RT.HYPERLINK for r in d.part.rels.values())
    question_tables=[t for t in d.tables if len(t.columns)==6]
    assert len(question_tables)==15
    assert sum(len(t.rows)-1 for t in question_tables)==100
    for q,row in zip(data['rows'],[row for t in question_tables for row in t.rows[1:]]):
        assert row.cells[0].text==q['id']+('*' if q['critical'] else '')
        assert row.cells[2].text==q['question']
        assert row.cells[3].text==str(q['weight']).replace('.0','').replace('.',',')+'%'
    print(json.dumps({'questions':100,'groups':15,'weight':100,'links':0,'file':str(dest)},ensure_ascii=False))

if __name__=='__main__':main()
