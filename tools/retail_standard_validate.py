"""Validate the generic retail package and render Word PDF pages for visual QA."""
from pathlib import Path
from urllib.parse import unquote
import re,json,hashlib,zipfile
import fitz
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/"wiki/requirements/retail_convenience_standard"
QA=ROOT/"tmp/retail-standard-qa"
def main():
    result={"checks":[],"rendered":[]}
    for prefix,standard,questionnaire,matrix,n in [
        ("RW","01_wms_standard","04_questionnaire_wms","06_vendor_comparison_wms",60),
        ("RT","02_tms_standard","05_questionnaire_tms","07_vendor_comparison_tms",40)]:
        a=(SRC/(standard+".md")).read_text(encoding="utf-8")
        b=(SRC/(questionnaire+".md")).read_text(encoding="utf-8")
        ids=re.findall(r"^### ("+prefix+r"\d\d) ",a,re.M)
        assert len(ids)==n and len(set(ids))==n,(prefix,ids)
        assert re.findall(r"^### ("+prefix+r"\d\d) ",b,re.M)==ids
        wa=list(map(int,re.findall(r"Профиль [ЯУР] · вес (\d+)%",a)))
        wb=list(map(int,re.findall(r"Вес (\d+)% · профиль",b)))
        assert len(wa)==n and sum(wa)==100 and wa==wb,(prefix,wa,wb)
        rows=[line for line in (SRC/(matrix+".md")).read_text(encoding="utf-8").splitlines() if re.match(r"^\| "+prefix+r"\d\d ",line)]
        assert len(rows)==n
        for row,ident,weight in zip(rows,ids,wa):
            cells=[c.strip() for c in row.strip("|").split("|")]
            assert len(cells)==12 and cells[0]==ident and int(cells[2])==weight,cells
            assert all(re.fullmatch(r"(Д|Ч) S\d\d|НД",c) for c in cells[3:])
        result["checks"].append({"catalog":prefix,"count":n,"weight":sum(wa),"vendor_answers":n*9})
    files=[p for p in SRC.glob("*.md") if p.name!="publication.md"]
    assert len(files)==11
    for p in files:
        text=p.read_text(encoding="utf-8")
        assert "<!--" not in text,p
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)",text):
            if not re.match(r"^[a-z]+:",target):
                assert (p.parent/unquote(target.split("#")[0])).exists(),(p.name,target)
        dp=SRC/"docx"/(p.stem+".docx")
        with zipfile.ZipFile(dp) as z:assert z.testzip() is None
        d=Document(dp)
        assert sum(x.style.name=="Title" for x in d.paragraphs)==1
        external={r.target_ref for r in d.part.rels.values() if r.reltype==RT.HYPERLINK and str(r.target_ref).startswith("http")}
        expected={u for u in re.findall(r"\[[^\]]+\]\((https?://[^)]+)\)",text)}
        assert expected==external,(p.name,expected-external,external-expected)
        for style in ["Title","Heading 1","Heading 2","Heading 3"]:
            assert str(d.styles[style].font.color.rgb)=="000000"
        for t in d.tables:
            assert t.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
            assert t._tbl.tblPr.xpath("./w:tblBorders")
    result["checks"].append({"files":len(files),"docx":len(files),"links":"preserved"})
    gap=(SRC/"08_current_system_gap.md").read_text(encoding="utf-8")
    headings="\n".join(re.findall(r"^### G\d+ .*",gap,re.M))
    all_ids=re.findall(r"R[WT]\d\d",headings)
    assert len(all_ids)==100 and len(set(all_ids))==100
    sourcefiles=[]
    for rel in re.findall(r"\]\(\.\./\.\./\.\./([^)#]+)#L\d+\)",gap):
        fp=ROOT/unquote(rel)
        sourcefiles.append({"path":str(fp.relative_to(ROOT)).replace("\\","/"),"sha256":hashlib.sha256(fp.read_bytes()).hexdigest()})
    result["source_baseline"]=sourcefiles
    expected_stems={p.stem for p in files}
    rendered_stems={p.stem for p in QA.glob("*.pdf") if p.stem in expected_stems}
    if rendered_stems:assert rendered_stems==expected_stems,("Incomplete PDF render",expected_stems-rendered_stems)
    for pdf in sorted(QA.glob("*.pdf")):
        if pdf.stem not in expected_stems:continue
        assert pdf.stat().st_mtime >= (SRC/"docx"/(pdf.stem+".docx")).stat().st_mtime,("Stale PDF",pdf)
        folder=QA/pdf.stem;folder.mkdir(exist_ok=True)
        doc=fitz.open(pdf)
        for i,page in enumerate(doc):
            pix=page.get_pixmap(dpi=105,alpha=False);pix.save(folder/f"page-{i+1:03}.png")
            words=page.get_text("words")
            outside=[w[:5] for w in words if w[0]<15 or w[1]<10 or w[2]>page.rect.width-15 or w[3]>page.rect.height-10]
            assert words and not outside,(pdf.name,i+1,outside)
            result["rendered"].append({"file":pdf.stem,"page":i+1,"words":len(words),"outside":outside,"image":str(folder/f"page-{i+1:03}.png")})
    if not result["rendered"] and (QA/"word_render.json").exists():
        word=json.loads((QA/"word_render.json").read_text(encoding="utf-8-sig"))
        if isinstance(word,dict):word=[word]
        for item in word:
            stem=Path(item["name"]).stem
            images=sorted((QA/stem).glob("page-*.png"))
            assert len(images)==item["pages"],(stem,len(images),item["pages"])
            for i,img in enumerate(images):
                result["rendered"].append({"file":stem,"page":i+1,"renderer":item["renderer"],"outside":[],"image":str(img)})
    (QA/"validation.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"checks":result["checks"],"evidence_files":len(sourcefiles),"pages":len(result["rendered"]),"outside_boxes":sum(len(x["outside"]) for x in result["rendered"]),"layout_note":"Visual PNG review required in addition to structural checks"},ensure_ascii=False,indent=2))
if __name__=="__main__":main()
