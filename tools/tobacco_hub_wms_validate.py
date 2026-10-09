"""Validate IDs, weights, Word links and rendered pages of the tobacco profile."""
from pathlib import Path
from urllib.parse import unquote
import re,json,zipfile
import fitz
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'wiki/requirements/tobacco_regional_hubs'
QA=ROOT/'tmp/tobacco-hub-wms-qa'

def main():
    a=(SRC/'01_functional_requirements.md').read_text(encoding='utf-8')
    b=(SRC/'03_vendor_questionnaire.md').read_text(encoding='utf-8')
    c=(SRC/'04_weights_acceptance.md').read_text(encoding='utf-8')
    ids=re.findall(r'^### TH (\d\d) ',a,re.M)
    assert ids==[f'{n:02}' for n in range(1,53)]
    assert re.findall(r'^### TH (\d\d) ',b,re.M)==ids
    wa=list(map(int,re.findall(r'^Вес (\d+)%',a,re.M)))
    wb=list(map(int,re.findall(r'^Вес (\d+)%',b,re.M)))
    wc=list(map(int,re.findall(r'^\| TH-\d\d \| [^|]+ \| (\d+)%',c,re.M)))
    assert len(wa)==52 and wa==wb==wc and sum(wa)==100
    assert len(re.findall(r'^\| AT\d\d ',c,re.M))==18
    files=[p for p in sorted(SRC.glob('*.md')) if p.name!='publication.md']
    assert len(files)==6
    result={'requirements':52,'questions':52,'weight':100,'scenarios':18,'files':[],'rendered':[]}
    for p in files:
        text=p.read_text(encoding='utf-8')
        assert '<!--' not in text
        for target in re.findall(r'\[[^\]]+\]\(([^)]+)\)',text):
            if not re.match(r'^[a-z]+:',target) and target!='publication.md':
                assert (p.parent/unquote(target.split('#')[0])).exists(),(p,target)
        dp=SRC/'docx'/(p.stem+'.docx')
        with zipfile.ZipFile(dp) as z:assert z.testzip() is None
        d=Document(dp)
        assert sum(x.style.name=='Title' for x in d.paragraphs)==1
        expected=set(re.findall(r'\[[^\]]+\]\((https?://[^)]+)\)',text))
        actual={r.target_ref for r in d.part.rels.values() if r.reltype==RT.HYPERLINK and str(r.target_ref).startswith('http')}
        assert expected==actual,(p,expected-actual,actual-expected)
        for s in ['Title','Heading 1','Heading 2','Heading 3']:assert str(d.styles[s].font.color.rgb)=='000000'
        for t in d.tables:
            assert t.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
            assert t._tbl.tblPr.xpath('./w:tblBorders')
        pdf=QA/(p.stem+'.pdf')
        assert pdf.exists() and pdf.stat().st_mtime>=dp.stat().st_mtime,('Missing or stale PDF',p)
        folder=QA/p.stem
        folder.mkdir(exist_ok=True)
        rendered=fitz.open(pdf)
        for n,page in enumerate(rendered,1):
            words=page.get_text('words')
            outside=[w[:5] for w in words if w[0]<15 or w[1]<10 or w[2]>page.rect.width-15 or w[3]>page.rect.height-10]
            assert words and not outside,(p,n,outside)
            img=folder/f'page-{n:03}.png'
            page.get_pixmap(dpi=105,alpha=False).save(img)
            result['rendered'].append({'file':p.stem,'page':n,'image':str(img),'words':len(words)})
        result['files'].append({'name':p.name,'pages':len(rendered),'bytes':dp.stat().st_size})
    (QA/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k!='rendered'},ensure_ascii=False,indent=2))
    print('All pages require visual review:',len(result['rendered']))

if __name__=='__main__':main()
